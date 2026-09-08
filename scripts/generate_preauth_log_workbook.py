"""
One-off script (not part of any Docker image / running service) - generates
an Excel workbook documenting the preauth_request_log and preauth_response_log
table structure, with real sample data pulled from the live database.

Usage:
    pip install -r scripts/requirements.txt
    python3 scripts/generate_preauth_log_workbook.py [output_path]

Connects to the same MySQL instance every service in this repo uses, via the
host-mapped port (localhost:3306) - run this from the host, not inside a
container.
"""
import sys
import json
import pymysql
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DB_CONFIG = {
    "host": "localhost",
    "port": 3306,
    "user": "fhir_user",
    "password": "fhir_password",
    "database": "fhir_db",
    "cursorclass": pymysql.cursors.DictCursor,
}

TABLES = ["preauth_request_log", "preauth_response_log"]

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
SECTION_FONT = Font(bold=True, size=13)


def fetch_structure(cursor, table_name):
    cursor.execute("""
        SELECT COLUMN_NAME, COLUMN_TYPE, IS_NULLABLE, COLUMN_KEY, COLUMN_DEFAULT, EXTRA, COLUMN_COMMENT
        FROM INFORMATION_SCHEMA.COLUMNS
        WHERE TABLE_SCHEMA = %s AND TABLE_NAME = %s
        ORDER BY ORDINAL_POSITION
    """, (DB_CONFIG["database"], table_name))
    return cursor.fetchall()


def fetch_sample_rows(cursor, table_name, limit=10):
    cursor.execute(f"SELECT * FROM {table_name} ORDER BY created_at DESC, id DESC LIMIT %s", (limit,))
    return cursor.fetchall()


def autosize_columns(ws, min_width=10, max_width=80):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            length = len(str(cell.value))
            col = cell.column_letter
            widths[col] = max(widths.get(col, 0), min(length + 2, max_width))
    for col, width in widths.items():
        ws.column_dimensions[col].width = max(width, min_width)


def write_table_sheet(wb, table_name, structure_rows, sample_rows):
    ws = wb.create_sheet(title=table_name)

    ws.cell(row=1, column=1, value=f"Table structure: {table_name}").font = SECTION_FONT
    struct_headers = ["Column Name", "Data Type", "Nullable", "Key", "Default", "Extra", "Comment"]
    header_row = 3
    for col_idx, header in enumerate(struct_headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    r = header_row + 1
    for col in structure_rows:
        ws.cell(row=r, column=1, value=col["COLUMN_NAME"])
        ws.cell(row=r, column=2, value=col["COLUMN_TYPE"])
        ws.cell(row=r, column=3, value=col["IS_NULLABLE"])
        ws.cell(row=r, column=4, value=col["COLUMN_KEY"])
        ws.cell(row=r, column=5, value=str(col["COLUMN_DEFAULT"]) if col["COLUMN_DEFAULT"] is not None else "")
        ws.cell(row=r, column=6, value=col["EXTRA"])
        ws.cell(row=r, column=7, value=col["COLUMN_COMMENT"])
        r += 1

    r += 2
    ws.cell(row=r, column=1, value=f"Sample data ({len(sample_rows)} most recent rows)").font = SECTION_FONT
    r += 2

    if sample_rows:
        sample_headers = list(sample_rows[0].keys())
        for col_idx, header in enumerate(sample_headers, start=1):
            cell = ws.cell(row=r, column=col_idx, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
        r += 1

        for row in sample_rows:
            for col_idx, header in enumerate(sample_headers, start=1):
                value = row[header]
                if header == "payload":
                    # pymysql does NOT auto-decode a native JSON column - it
                    # comes back as the raw JSON text already, not a Python
                    # dict/list. Parse + re-dump only to pretty-print it;
                    # json.dumps()-ing the raw string directly here would
                    # double-encode it (wrap already-valid JSON text in an
                    # extra layer of string escaping).
                    try:
                        value = json.dumps(json.loads(value), indent=2, default=str)
                    except (TypeError, ValueError):
                        pass
                elif value is not None and not isinstance(value, (str, int, float)):
                    value = str(value)
                ws.cell(row=r, column=col_idx, value=value)
            r += 1
    else:
        ws.cell(row=r, column=1, value="(no rows yet - run the PreAuth flow first, e.g. POST /preauth/{claim_id})")

    autosize_columns(ws)


def main():
    output_path = sys.argv[1] if len(sys.argv) > 1 else "preauth_log_tables.xlsx"

    conn = pymysql.connect(**DB_CONFIG)
    try:
        wb = Workbook()
        wb.remove(wb.active)  # drop the default blank sheet

        with conn.cursor() as cursor:
            for table_name in TABLES:
                structure_rows = fetch_structure(cursor, table_name)
                sample_rows = fetch_sample_rows(cursor, table_name)
                write_table_sheet(wb, table_name, structure_rows, sample_rows)
                print(f"✅ {table_name}: {len(structure_rows)} columns, {len(sample_rows)} sample rows")

        wb.save(output_path)
        print(f"✅ Workbook written to {output_path}")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
