"""
Workflow Service - FastAPI
CRUD for saved workflow graphs, plus the execution engine and run history API.
"""
import logging
import json
import asyncio
from datetime import datetime
from io import BytesIO
from fastapi import FastAPI, HTTPException, Depends
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import sys
import os

# Add shared module to path
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '../../'))

from shared import (
    get_db, create_kafka_producer, TOPICS,
    Workflow, WorkflowRun, WorkflowRunStep, Base, engine
)

from .schemas import (
    WorkflowIn, WorkflowSummary, WorkflowOut, RunRequest, RunAccepted,
    WorkflowRunOut, WorkflowRunStepOut, WorkflowRunSummary
)
from .node_registry import NODE_TYPES
from .engine import run_workflow, validate_graph, GraphValidationError
from .executors import ExecutionContext

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Initialize FastAPI app
app = FastAPI(
    title="Workflow Service",
    description="n8n-style workflow builder for JSON2FHIR",
    version="1.0.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Initialize database tables
try:
    Base.metadata.create_all(bind=engine)
    logger.info("✅ Database tables initialized")
except Exception as e:
    logger.error(f"❌ Error initializing database: {e}")

# Global producer / execution context
producer = None
ctx: ExecutionContext = None


@app.on_event("startup")
async def startup_event():
    global producer, ctx
    try:
        producer = await create_kafka_producer()
        ctx = ExecutionContext(producer)
        logger.info("🚀 Workflow Service started")
    except Exception as e:
        logger.error(f"❌ Failed to start producer: {e}")
        raise


@app.on_event("shutdown")
async def shutdown_event():
    global producer
    if producer:
        await producer.stop()
        logger.info("🛑 Workflow Service stopped")


@app.get("/health", tags=["Health"])
async def health_check():
    return {"status": "healthy", "service": "Workflow Service", "version": "1.0.0"}


@app.get("/node-types", tags=["Registry"])
async def get_node_types():
    return {"node_types": NODE_TYPES}


@app.get("/topics", tags=["Registry"])
async def get_topics():
    return {"topics": TOPICS}


def _workflow_to_out(w: Workflow) -> WorkflowOut:
    return WorkflowOut(
        id=w.id, name=w.name, description=w.description,
        definition=json.loads(w.definition),
        created_at=w.created_at, updated_at=w.updated_at
    )


@app.post("/workflows", response_model=WorkflowOut, status_code=201, tags=["Workflows"])
async def create_workflow(body: WorkflowIn, db: Session = Depends(get_db)):
    workflow = Workflow(
        name=body.name,
        description=body.description,
        definition=json.dumps(body.definition)
    )
    db.add(workflow)
    db.commit()
    db.refresh(workflow)
    logger.info(f"✅ Created workflow: {workflow.id} ({workflow.name})")
    return _workflow_to_out(workflow)


@app.get("/workflows", tags=["Workflows"])
async def list_workflows(limit: int = 50, offset: int = 0, db: Session = Depends(get_db)):
    rows = db.query(Workflow).order_by(Workflow.updated_at.desc()).offset(offset).limit(limit).all()
    return {"workflows": [WorkflowSummary.model_validate(w) for w in rows]}


@app.get("/workflows/{workflow_id}", response_model=WorkflowOut, tags=["Workflows"])
async def get_workflow(workflow_id: int, db: Session = Depends(get_db)):
    workflow = db.query(Workflow).filter(Workflow.id == workflow_id).first()
    if not workflow:
        raise HTTPException(status_code=404, detail=f"Workflow {workflow_id} not found")
    return _workflow_to_out(workflow)


@app.put("/workflows/{workflow_id}", response_model=WorkflowOut, tags=["Workflows"])
async def update_workflow(workflow_id: int, body: WorkflowIn, db: Session = Depends(get_db)):
    workflow = db.query(Workflow).filter(Workflow.id == workflow_id).first()
    if not workflow:
        raise HTTPException(status_code=404, detail=f"Workflow {workflow_id} not found")
    workflow.name = body.name
    workflow.description = body.description
    workflow.definition = json.dumps(body.definition)
    db.commit()
    db.refresh(workflow)
    logger.info(f"✅ Updated workflow: {workflow.id}")
    return _workflow_to_out(workflow)


@app.delete("/workflows/{workflow_id}", status_code=204, tags=["Workflows"])
async def delete_workflow(workflow_id: int, db: Session = Depends(get_db)):
    workflow = db.query(Workflow).filter(Workflow.id == workflow_id).first()
    if not workflow:
        raise HTTPException(status_code=404, detail=f"Workflow {workflow_id} not found")
    try:
        db.delete(workflow)
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Cannot delete a workflow with run history"
        )
    logger.info(f"✅ Deleted workflow: {workflow_id}")


@app.post("/workflows/{workflow_id}/run", response_model=RunAccepted, status_code=202, tags=["Runs"])
async def run_workflow_endpoint(workflow_id: int, body: RunRequest, db: Session = Depends(get_db)):
    workflow = db.query(Workflow).filter(Workflow.id == workflow_id).first()
    if not workflow:
        raise HTTPException(status_code=404, detail=f"Workflow {workflow_id} not found")

    definition = json.loads(workflow.definition)

    try:
        validate_graph(definition)
    except GraphValidationError as e:
        raise HTTPException(status_code=400, detail=str(e))

    trigger_input = body.trigger_input or {}

    run = WorkflowRun(
        workflow_id=workflow_id, status="RUNNING",
        trigger_input=json.dumps(trigger_input), started_at=datetime.utcnow()
    )
    db.add(run)
    db.commit()
    db.refresh(run)

    logger.info(f"🚀 Starting run {run.id} for workflow {workflow_id}")
    asyncio.create_task(run_workflow(run.id, definition, trigger_input, ctx))

    return RunAccepted(
        run_id=run.id, workflow_id=workflow_id, status=run.status, started_at=run.started_at
    )


@app.get("/runs/{run_id}", response_model=WorkflowRunOut, tags=["Runs"])
async def get_run(run_id: int, db: Session = Depends(get_db)):
    run = db.query(WorkflowRun).filter(WorkflowRun.id == run_id).first()
    if not run:
        raise HTTPException(status_code=404, detail=f"Run {run_id} not found")

    steps = db.query(WorkflowRunStep).filter(WorkflowRunStep.run_id == run_id).order_by(WorkflowRunStep.id).all()

    return WorkflowRunOut(
        id=run.id, workflow_id=run.workflow_id, status=run.status,
        trigger_input=json.loads(run.trigger_input) if run.trigger_input else None,
        error=run.error, started_at=run.started_at, finished_at=run.finished_at,
        steps=[
            WorkflowRunStepOut(
                node_id=s.node_id, node_type=s.node_type, status=s.status,
                input=json.loads(s.input) if s.input else None,
                output=json.loads(s.output) if s.output else None,
                error=s.error, started_at=s.started_at, finished_at=s.finished_at
            )
            for s in steps
        ]
    )


EXCEL_HEADER_FONT = Font(bold=True, color="FFFFFF")
EXCEL_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
EXCEL_SECTION_FONT = Font(bold=True, size=13)

# Labels for the two payloads a transform node's step recorded - which one
# is "before" and which is "after" the transform, in plain terms.
CHAIN_STEP_LABELS = {
    "json_to_fhir": ("JSON Request", "FHIR Request"),
    "fhir_to_json": ("FHIR Response", "JSON Response"),
    "fhir_transform": ("Input", "Output"),  # legacy node type, direction ambiguous
}
TRANSFORM_NODE_TYPES = set(CHAIN_STEP_LABELS.keys())


def _find_nearest_transform_ancestors(node_id: str, node_by_id: dict, incoming: dict) -> list:
    """
    BFS backwards from node_id through pass-through hops (Kafka Publish, DB
    Stored Procedure, View, ...) to find the nearest JSON<->FHIR transform
    node(s) feeding it. An Excel node wired a few steps downstream of a
    transform (e.g. after the Kafka Publish that follows it) should still
    show that transform's JSON/FHIR pair, not the pass-through node's
    identical input/output.
    """
    visited = set()
    frontier = list(incoming.get(node_id, []))
    found = []
    while frontier:
        nid = frontier.pop(0)
        if nid in visited:
            continue
        visited.add(nid)
        if node_by_id.get(nid, {}).get("type") in TRANSFORM_NODE_TYPES:
            found.append(nid)
        else:
            frontier.extend(incoming.get(nid, []))
    return found


@app.get("/runs/{run_id}/nodes/{node_id}/export/excel", tags=["Runs"])
async def export_run_node_excel(run_id: int, node_id: str, db: Session = Depends(get_db)):
    """
    Downloads an Excel workbook of whatever actually flowed through this run
    at the nearest JSON<->FHIR transform node upstream of node_id (walking
    back through pass-through hops if needed) - its own recorded input and
    output for this run, e.g. the JSON that went into a JSON->FHIR node and
    the FHIR that came out. Meant for the Excel Export node's Download
    button: no claim ID or PreAuth DB lookup involved, purely this run's own
    step history (workflow_run_steps).
    """
    run = db.query(WorkflowRun).filter(WorkflowRun.id == run_id).first()
    if not run:
        raise HTTPException(status_code=404, detail=f"Run {run_id} not found")

    workflow = db.query(Workflow).filter(Workflow.id == run.workflow_id).first()
    if not workflow:
        raise HTTPException(status_code=404, detail=f"Workflow {run.workflow_id} not found")

    definition = json.loads(workflow.definition)
    node_by_id = {n["id"]: n for n in definition.get("nodes", []) or []}
    if node_id not in node_by_id:
        raise HTTPException(status_code=404, detail=f"Node '{node_id}' not found in workflow {run.workflow_id}")

    incoming: dict = {}
    for e in definition.get("edges", []) or []:
        incoming.setdefault(e.get("target"), []).append(e.get("source"))

    if not incoming.get(node_id):
        raise HTTPException(
            status_code=400,
            detail=f"Node '{node_id}' has no upstream node - wire it after a JSON → FHIR or FHIR → JSON node first"
        )

    # Prefer the nearest actual transform node upstream; if there isn't one
    # anywhere in this branch (e.g. wired straight after the trigger), fall
    # back to the immediate predecessor(s) so the button still shows *something*.
    source_ids = _find_nearest_transform_ancestors(node_id, node_by_id, incoming) or incoming[node_id]

    steps = db.query(WorkflowRunStep).filter(
        WorkflowRunStep.run_id == run_id,
        WorkflowRunStep.node_id.in_(source_ids)
    ).all()
    steps_by_node = {s.node_id: s for s in steps}

    wb = Workbook()
    wb.remove(wb.active)

    for pred_id in source_ids:
        step = steps_by_node.get(pred_id)
        pred_type = node_by_id.get(pred_id, {}).get("type", "")
        in_label, out_label = CHAIN_STEP_LABELS.get(pred_type, ("Input", "Output"))

        ws = wb.create_sheet(title=pred_id[:31] or "sheet")
        ws.cell(row=1, column=1, value=f"Source node: {pred_id} ({pred_type or 'unknown type'})").font = EXCEL_SECTION_FONT
        ws.cell(row=2, column=1, value=f"run_id: {run_id}, status: {step.status if step else 'no data for this run'}")

        header_row = 4
        for col_idx, header in enumerate(["Field", "Payload"], start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=header)
            cell.font = EXCEL_HEADER_FONT
            cell.fill = EXCEL_HEADER_FILL

        r = header_row + 1
        if step:
            for label, raw in ((in_label, step.input), (out_label, step.output)):
                ws.cell(row=r, column=1, value=label)
                try:
                    payload = json.dumps(json.loads(raw), indent=2, default=str) if raw else "(none)"
                except (TypeError, ValueError):
                    payload = raw
                payload_cell = ws.cell(row=r, column=2, value=payload)
                payload_cell.alignment = Alignment(wrap_text=True, vertical="top")
                r += 1
            if step.error:
                ws.cell(row=r, column=1, value="Error")
                ws.cell(row=r, column=2, value=step.error)
        else:
            ws.cell(row=r, column=1, value="(this node didn't run in this run)")

        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 100

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=run_{run_id}_{node_id}_chain.xlsx"}
    )


@app.get("/workflows/{workflow_id}/runs", tags=["Runs"])
async def list_runs(workflow_id: int, db: Session = Depends(get_db)):
    workflow = db.query(Workflow).filter(Workflow.id == workflow_id).first()
    if not workflow:
        raise HTTPException(status_code=404, detail=f"Workflow {workflow_id} not found")
    runs = db.query(WorkflowRun).filter(WorkflowRun.workflow_id == workflow_id).order_by(WorkflowRun.started_at.desc()).all()
    return {"runs": [WorkflowRunSummary.model_validate(r) for r in runs]}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(
        app,
        host="0.0.0.0",
        port=8002,
        log_level="info"
    )
