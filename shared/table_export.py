"""
Excel workbooks that show stored data the way the database holds it: one
sheet per table, laid out as a real table (header row of column names, then
the rows) - nothing else.

What is exported depends on the endpoint that stored the data: PROFILES
lists, per endpoint, the tables it writes to and how to pick that record's
rows inside each of them.
"""
import json
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text

from .preauth_tables import TABLE_ORDER, ensure_preauth_schema

EXCEL_CELL_LIMIT = 32767  # Excel rejects longer cell text (base64 attachments, big JSON payloads)
MAX_COLUMN_WIDTH = 60

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="305496")

_CLAIM_FILTERS = {
    "message_header": "id IN (SELECT message_header_id FROM claim WHERE id = :k)",
    "claim": "id = :k",
    "claim_related": "claim_id = :k",
    "claim_request_insurance": "claim_id = :k",
    "diagnosis": "claim_id = :k",
    "care_team": "claim_id = :k",
    "supporting_info": "claim_id = :k",
    "claim_request_item": "claim_id = :k",
    "claim_request_detail": "item_id IN (SELECT id FROM claim_request_item WHERE claim_id = :k)",
    "coverage_class": "insurance_id IN (SELECT id FROM claim_request_insurance WHERE claim_id = :k)",
    "claim_request_encounter": "identifier IN (SELECT encounter_identifier FROM claim WHERE id = :k)",
}
_TRACKING_FILTERS = {"message_tracking": "correlation_id = :k", "audit_log": "correlation_id = :k"}
_TRACKING_ORDER = {"message_tracking": "created_at DESC", "audit_log": "id DESC"}

# endpoint -> the tables it stores into. The eligibility and patient endpoints
# only write the tracking/audit tables; PreAuth claims write the 11 claim tables.
PROFILES: Dict[str, Dict[str, Any]] = {
    "preauth-claim": {
        "label": "PreAuth claim",
        "endpoint": "POST /api/v1/preauth/responses (communication-service)",
        "key_name": "claim_id",
        "tables": list(TABLE_ORDER),
        "filters": _CLAIM_FILTERS,
        "latest_order": {},
        "all_rows_limit": None,
    },
    "eligibility-request": {
        "label": "Eligibility request",
        "endpoint": "POST /api/v1/eligibility/requests (integration-api)",
        "key_name": "correlation_id",
        "tables": ["message_tracking", "audit_log"],
        "filters": _TRACKING_FILTERS,
        "latest_order": _TRACKING_ORDER,
        "all_rows_limit": 200,
    },
    "eligibility-response": {
        "label": "Eligibility response",
        "endpoint": "POST /api/v1/eligibility/responses (communication-service)",
        "key_name": "correlation_id",
        "tables": ["message_tracking", "audit_log"],
        "filters": _TRACKING_FILTERS,
        "latest_order": _TRACKING_ORDER,
        "all_rows_limit": 200,
    },
    "patient": {
        "label": "Patient",
        "endpoint": "POST /api/v1/patients (integration-api)",
        "key_name": "correlation_id",
        "tables": ["message_tracking", "audit_log"],
        "filters": _TRACKING_FILTERS,
        "latest_order": _TRACKING_ORDER,
        "all_rows_limit": 200,
    },
}


def _cell_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, bytes):
        value = value.decode("utf-8", errors="replace")
    if isinstance(value, (dict, list)):
        value = json.dumps(value, default=str)
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        if len(value) > EXCEL_CELL_LIMIT:
            keep = EXCEL_CELL_LIMIT - 40
            value = f"{value[:keep]}...[truncated, {len(value)} chars]"
    return value


def _style_header(sheet, row: int = 1) -> None:
    for cell in sheet[row]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center")


def _fit_columns(sheet, widths: List[int]) -> None:
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = min(width + 2, MAX_COLUMN_WIDTH)


def _write_data_sheet(wb: Workbook, conn, quote, name: str, where: Optional[str], key: Optional[str],
                      order_by: Optional[str], limit: Optional[int]) -> int:
    """One table as a sheet: header = the table's own columns, then its rows. Returns rows written."""
    sql = f"SELECT * FROM {quote(name)}"
    if where:
        sql += f" WHERE {where}"
    sql += f" ORDER BY {order_by or '1'}"
    if limit:
        sql += f" LIMIT {int(limit)}"
    result = conn.execute(text(sql), {"k": key} if where else {})

    columns = list(result.keys())
    sheet = wb.create_sheet(title=name)
    sheet.append(columns)
    widths = [len(c) for c in columns]
    count = 0
    for row in result:
        values = [_cell_value(v) for v in row]
        sheet.append(values)
        count += 1
        for i, v in enumerate(values):
            if isinstance(v, (datetime, date)):
                sheet.cell(row=sheet.max_row, column=i + 1).number_format = "yyyy-mm-dd hh:mm:ss"
            widths[i] = max(widths[i], len(str(v)) if v is not None else 0)

    _style_header(sheet)
    _fit_columns(sheet, widths)
    sheet.freeze_panes = "A2"
    if columns:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(count + 1, 1)}"
    return count


def build_endpoint_workbook(engine, profile_key: str, key: Optional[str] = None, scope: str = "this-run") -> Workbook:
    """
    Workbook of the tables the given endpoint stores into, one sheet each.
    Scope "this-run" shows only the rows of the record identified by key
    (claim_id / correlation_id) - no rows at all if there is no such record;
    "all-rows" shows every stored row (tracking tables: the latest 200).
    """
    profile = PROFILES[profile_key]
    if profile_key == "preauth-claim":
        ensure_preauth_schema(engine)  # creates any missing table before we read it

    scoped = scope != "all-rows"

    wb = Workbook()
    wb.remove(wb.active)

    quote = engine.dialect.identifier_preparer.quote
    with engine.connect() as conn:
        for name in profile["tables"]:
            if scoped:
                where, limit = profile["filters"][name], None
                order = "id" if name == "audit_log" else "1"  # audit trail in the order it happened
            else:
                where, order, limit = None, profile["latest_order"].get(name), profile["all_rows_limit"]
            _write_data_sheet(wb, conn, quote, name, where, key, order, limit)
    return wb


def workbook_bytes(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
